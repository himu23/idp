"""Build the final presentation (report/Prithvi_Flood-IN_Final.pptx) from the analysis outputs.

    python report/build_slides.py
Numbers are read from results/analysis/summary.json etc., figures from results/analysis/figures,
pinn/results/figures and mahanadi/outputs, so the slides always match the report.
"""
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
A = os.path.join(ROOT, "results", "analysis")
S = json.load(open(os.path.join(A, "summary.json")))
MAH = json.load(open(os.path.join(ROOT, "mahanadi", "outputs", "summary.json")))
INK, MUTED, ACCENT = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x2A, 0x78, 0xD6)


def get(proto, model, mode="full", mods="default"):
    return next((s for s in S if s["protocol"] == proto and s["model"] == model and s["mode"] == mode
                 and s["modalities"] == mods), None)


def iou(proto, model, mode="full", mods="default", sd=True):
    s = get(proto, model, mode, mods) or get(proto, model, "none", mods)
    if s is None:
        return "–"
    return f"{s['test_iou']:.3f} ± {s['test_iou_std']:.3f}" if sd and s["n_seeds"] > 1 else f"{s['test_iou']:.3f}"


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, lines, size=18, color=INK, bold_first=False, bullet=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " if bullet and line else "") + line
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold_first and i == 0
        p.space_after = Pt(6)
    return tb


def slide(title, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    text(s, 0.5, 0.3, 12.3, 0.8, [title], size=30, bold_first=True)
    if subtitle:
        text(s, 0.5, 1.0, 12.3, 0.6, [subtitle], size=16, color=MUTED)
    text(s, 0.5, 7.0, 12.3, 0.4, ["Prithvi Flood-IN · CE 491 IDP · IIT Bombay × ClimMaTech · final review"], size=10,
         color=MUTED)
    return s


def picture(s, path, x, y, w=None, h=None):
    if not os.path.exists(path):
        text(s, x, y, 4, 0.5, [f"[missing figure: {os.path.basename(path)}]"], size=12, color=MUTED)
        return
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    s.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def table(s, x, y, w, rows, col_w=None, size=12):
    shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.35 * len(rows)))
    t = shp.table
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = str(v)
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(size)
                p.font.bold = i == 0
    if col_w:
        for j, cw in enumerate(col_w):
            t.columns[j].width = Inches(cw)
    return t


exec(open(os.path.join(HERE, "slides_content.py"), encoding="utf-8").read())
out = os.path.join(HERE, "Prithvi_Flood-IN_Final.pptx")
prs.save(out)
print("written", out, len(prs.slides), "slides")
