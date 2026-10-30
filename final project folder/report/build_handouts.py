"""Hand-outs built from the same result files as the report, so their numbers always match it.

    python report/build_handouts.py          (after build_slides.py)

Writes in docs/ (Markdown and Word):
  EXECUTIVE_SUMMARY   two pages for the evaluators, the guide and the ClimMaTech mentors
  SLIDES_TEXT         every slide: what it shows, who presents it and what to say (the same notes go into the
                      speaker notes of report/Prithvi_Flood-IN_Final.pptx)
  TEAM_GUIDE          the whole project from the basics, for a team member who was not part of the work
"""
import glob
import json
import os

import pandas as pd
import pypandoc
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
AN, AT, IM = (os.path.join(ROOT, "results", d) for d in ("analysis", "atlas", "impact"))
DOCS = os.path.join(ROOT, "docs")
os.makedirs(DOCS, exist_ok=True)
DECK = os.path.join(HERE, "Prithvi_Flood-IN_Final.pptx")


def J(*p):
    path = os.path.join(*p)
    return json.load(open(path)) if os.path.exists(path) else {}


def n(x):
    return f"{float(x):,.0f}"


BENCH = {f"{s['protocol']}/{s['model']}/{s['mode']}/{s['modalities']}": s for s in J(AN, "summary.json")}
u = lambda k: BENCH[k]["test_iou"]  # noqa: E731

def top_floods(drv, k=2, gap_days=10):
    """The k largest floods on radar that are separate events (more than gap_days apart), largest first."""
    d = drv.assign(when=pd.to_datetime(drv.date)).sort_values("flood_km2", ascending=False)
    picked = []
    for _, r_ in d.iterrows():
        if all(abs((r_['when'] - p_['when']).days) > gap_days for p_ in picked):
            picked.append(r_)
        if len(picked) == k:
            break
    return picked


def fold_word(x):
    k_ = round(x)
    return {2: "two-fold", 3: "three-fold", 4: "four-fold", 5: "five-fold", 6: "six-fold"}.get(k_, f"{k_}-fold")


AS, R20, R20N, AUD = J(AT, "atlas_summary.json"), J(AN, "radar20m.json"), J(AN, "radar20m_r20.json"), J(AN, "audit_training.json")
# the atlas was built with the radar pair retrained on 20 m input when its threshold is that pair's validation threshold
ADOPTED = bool(R20N) and abs(AS.get("threshold", 0) - R20N.get("threshold", -1)) < 1e-9
PROD = R20N if ADOPTED else R20

LRN, WS, S2C, TT = J(AT, "learned", "learned_summary.json"), J(IM, "watershed_summary.json"), J(IM, "s2_crosscheck.json"), J(IM, "travel_time.json")
STR, GRD, OCN, CMP, SIT = (J(IM, f) for f in ("structures_summary.json", "ground_summary.json", "ocean_summary.json",
                                             "cm_plan.json", "siting_summary.json"))
EXS = J(IM, "exposure_summary.json")
EX = pd.read_csv(os.path.join(IM, "exposure_scenarios.csv")).set_index("scenario")
TR = pd.read_csv(os.path.join(IM, "transport_scenarios.csv")).set_index("scenario")
T25, TR25 = EX.loc["T25"], TR.loc["T25"]
V = AS["validation"]
HF, ALLP = V["high_flow_passes"], V["all_monsoon_passes"]
RP = {int(k): v for k, v in AS["gumbel"]["return_period_m3s"].items()}
RB = {mm: J(IM, "bulletins", f"bulletin_{round(RP[5])}_rain{mm}.json") for mm in (20, int(AS["heavy_rain_scenario"]["rain_mm"]))}
RH = int(AS["heavy_rain_scenario"]["rain_mm"])
B25 = J(IM, "bulletins", f"bulletin_{round(RP[25])}.json")
LI, LS, LA = LRN["pixel_iou"], LRN["settlements"], LRN["ranking_auc"]
TOPS = top_floods(pd.read_csv(os.path.join(IM, "flood_drivers.csv")))
PRED_WORDS = {"Q_same_day": "the flow of the same day", "Q_prev_day": "the flow of the previous day",
              "Q_max_3d": "the highest daily flow of the last three days", "Q_max_5d": "the highest daily flow of the last five days",
              "Q_max_7d": "the highest daily flow of the last seven days"}.get(AS["predictor"], "the recent flow")
SC = LRN.get("scenarios_all_years_model", {})
wT25 = STR["walls"]["T25"]
dcrit = STR["walls"]["critical_depth_m"]
fe = GRD["foundation_example"]
g25 = GRD["scenarios"]["T25"]
lr = CMP["link_raising"]
_ns = len(lr.get("links_needing_structure", []))
STRUCT = (f"; {_ns} more need a bridge or vented causeway" if _ns else
          f", all as road embankments (none needs more than {lr.get('max_raise_for_embankment_m', 3):.0f} m)")
dep = CMP["deployment"]
cnt = OCN["top20_flood_passes_driver_counts"]
s2a, s2f = S2C["all/all_months"], S2C["outside_normal/all_months"]
om = AUD["official_test_miou"]
sett_hf = V["settlements"]["high_flow_passes"]["model"]
auc = V["ranking_auc"]
cov15 = SIT["coverage"]["1500"]
M = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]
win = "December–April" if set(CMP["work_window_months"]) == {"Dec", "Jan", "Feb", "Mar", "Apr"} else ", ".join(CMP["work_window_months"])
scour25 = STR["bridges"]["pier_scour_T100_share25_range_m"]
FCS, RISKS = J(ROOT, "results", "forecast", "forecast_summary.json"), J(IM, "risk_economics.json")
LIVES, LCFG = J(ROOT, "live", "out", "mahanadi", "status.json"), J(ROOT, "live", "basins", "mahanadi.json")
BASK = sorted(d for d in os.listdir(os.path.join(ROOT, "results", "basins"))
              if os.path.exists(os.path.join(ROOT, "results", "basins", d, "atlas", "atlas_summary.json"))) \
    if os.path.isdir(os.path.join(ROOT, "results", "basins")) else []
BNAMES = [J(ROOT, "live", "basins", f"{k}.json").get("name", k) for k in BASK]
_bcfg = [J(ROOT, "live", "basins", f"{k}.json") for k in BASK]
LEVELS_TXT = ("; both run in the live feed, 'Be aware' from " + " and ".join(
    f"{n(c['level_flows_m3s']['1'])} m³/s at {c['gauge_name']}" for c in _bcfg if c.get("level_flows_m3s", {}).get("1"))
    ) if _bcfg and all(c.get("level_flows_m3s", {}).get("1") for c in _bcfg) else ""
_rps = sorted(glob.glob(os.path.join(ROOT, "live", "out", "replay_*", "mahanadi", "status.json")))
REPS = J(_rps[-1]) if _rps else {}
REPLAY_TXT = (f" `--replay {REPS['issued'][:10]}` shows what it would have said on that morning, before the largest flood of the "
              f"radar era: '{REPS['level_name']}' for a {n(REPS['peak']['median_m3s'])} m³/s peak on {REPS['peak']['date']}.") if REPS else ""
fk = FCS.get("skill", {})
EXEC_NEW = ""
if FCS and LIVES:
    EXEC_NEW += (f"7. **Forecasts and live warnings.** A model forecasts the Naraj flow one to three days ahead from the upstream "
                 f"flows and catchment rain (efficiency {fk['lead_3d']['boosted']['nse']:.2f} three days ahead, against "
                 f"{fk['lead_3d']['persistence']['nse']:.2f} for assuming no change). A live feed reads each day's GloFAS ensemble, "
                 "ECMWF rain and the official NDMA and CWC alerts and issues an impact-based warning: the villages to warn with "
                 "dates, the villages that lose their hospital, the roads that close, a CAP alert and an Atom feed.\n")
def _basin_facts(k):
    a = J(ROOT, "results", "basins", k, "atlas", "atlas_summary.json")
    sc = J(ROOT, "results", "basins", k, "impact", "exposure_summary.json").get("scenarios", {})
    ob = next((sc[s] for s in sc if s.startswith("observed_")), None)
    return (f"{a['n_passes']} passes, {n(a['flooded_ever_km2'])} km² flooded at least once"
            + (f", {n(ob['people'])} people on the largest flood seen" if ob else ""))


if BASK:
    EXEC_NEW += (f"8. **More basins.** The same pipeline mapped the "
                 + " and the ".join(f"{nm} ({_basin_facts(k)})" for k, nm in zip(BASK, BNAMES))
                 + " from every radar pass since 2017; both run in the live feed under the same impact rule.\n")
if RISKS:
    EXEC_NEW += (f"{7 + bool(FCS and LIVES) + bool(BASK)}. **Risk in rupees.** Flood damage to houses and paddy is at least "
                 f"₹{RISKS['expected_annual_damage_rs_crore']['total_rs_crore']:.0f} crore a year; raising the ten critical roads "
                 f"repays its cost about {RISKS['road_raising']['benefit_cost_ratio']:.0f} times in hospital access alone.\n")

# ---------------------------------------------------------------- executive summary
EXEC = f"""# Prithvi Flood-IN: executive summary

**What floods when the Mahanadi rises: a satellite flood atlas, flow-to-flood maps and engineering impacts for the
Mahanadi delta.** Integrated Design Project, Department of Civil Engineering, IIT Bombay, 2026. Team: Nikhil
Upadhyay (lead), Vineet Jangir, Prakhar Gupta, Himanshu Shete, Prabhat Kumar. Faculty guide: Prof. Basudev Biswal.
Industry partner: ClimMaTech (IIT Mandi), Dr. Vivek Gupta and Siddik Barbhuiya.

## The problem

Flood warning systems in India, including our partner's radar gauges, sirens and app, measure and forecast the river
at a point. None of them shows what a forecast flow floods: which fields, villages, roads, bridges and river banks go
under, and who must be warned. We built that missing step for the Mahanadi delta below Naraj (3,038 km², about
{EXS['population_ghsl'] / 1e6:.1f} million people), which floods every few years.

## What we built

1. **A fair model benchmark.** Six satellite foundation models (Prithvi-EO-2.0, TerraMind, DOFA, Clay, Satlas,
   DeCUR) against a CNN and a U-Net, 92 training runs, tested on an Indian flood none had seen. A U-Net won
   ({u('lio/unet/full/default'):.3f} water IoU against {u('lio/prithvi/full/default'):.3f} for the best foundation model;
   fusion {u('lio/enstop3_mean/full/default'):.3f}). An audit of all 92 runs found no problems and reproduces the published
   TerraMind scores within a point ({om['official_terramind_full_s0']['miou'] * 100:.1f} against 90.62).
2. **A flood atlas.** Monsoon floods are under cloud, so a radar-only TerraMind + U-Net model mapped all
   {AS['n_passes']} Sentinel-1 passes over the delta from 2017 to 2026: {n(AS['flooded_ever_km2'])} km² flooded at least once.
3. **Flow-to-flood maps.** Every 20 m pixel gets the Naraj flow at which it starts to flood, so a forecast flow
   becomes a map. On years held out from fitting it beats flood climatology at high flows (IoU
   {HF['model']['iou']:.2f} against {HF['climatology']['iou']:.2f}). A learned model that adds rain and flood history does
   better still on heavy-rain passes ({LI['heavy_rain_passes']['learned']['iou']:.2f} against {LI['heavy_rain_passes']['trigger']['iou']:.2f}).
4. **Impacts across the department's specializations.** The same flood scenarios answer one question for each of
   transportation, structural, geotechnical, ocean, construction management, water resources (with the watershed
   above the delta) and remote sensing.
5. **Tools for the partner.** A forecast-to-bulletin tool that also writes CAP alerts, an interactive map viewer,
   siren and river-gauge sites, and a project website.
6. **A field-visit kit** that measures what the models cannot know: accuracy at 40 random points, discharge and
   roughness, bridges, bank soils, well levels.
{EXEC_NEW}
## Key numbers (25-year flood, {n(RP[25])} m³/s at Naraj, unless stated)

| What | Number |
|---|---|
| Land flooded / people living on it | {T25['flood_km2']:.0f} km² / {n(T25['people'])} |
| Settlements with 10 % or more of people on flooded land | {int(T25['settlements >= 10 % of people flooded'])} |
| People cut off from every hospital (transportation) | {n(TR25['hospital: people cut off'])} in {int(TR25['hospital: settlements cut off'])} settlements |
| Extra travel every day | {n(TR25['town: extra person-hours of travel per day'])} person-hours |
| People where the flood is 1 m deep or more; brick walls crack at (structural) | {n(wT25['people_deeper_than_1m'])}; about {dcrit['230 mm wall, f_t 0.1 MPa']:.1f}–{dcrit['230 mm wall, f_t 0.2 MPa']:.1f} m |
| Pier scour below high flood level, a quarter of the 100-year flow | {scour25[0]:.0f}–{scour25[1]:.0f} m |
| People on soaked ground beside the flood; bearing loss (geotechnical) | {n(g25['people_in_zone'])}; about {fe['silty_sand_drained']['reduction_pct']} % |
| Largest floods that followed a Bay of Bengal storm (ocean) | {cnt['cyclone/depression']} of 20 |
| Work window; earthwork to raise critical roads (construction) | {win}; {n(lr['earthwork_m3'])} m³ (about ₹{lr['cost_crore_rs_at_350_per_m3']:.2f} crore) |
| Share of a flood peak from Hirakud dam (watershed) | median {WS['median_share_hirakud_of_naraj_peak']:.0%} |
| Sirens and coverage | {cov15['sirens']} sirens reach {cov15['share']:.0%} of exposed people at 1.5 km |
{f"| Expected flood damage a year (houses + paddy, lower bound) | ₹{RISKS['expected_annual_damage_rs_crore']['total_rs_crore']:.0f} crore |" + chr(10) + f"| Raising the ten critical roads: benefit / cost | {RISKS['road_raising']['benefit_cost_ratio']:.0f} : 1 |" + chr(10) if RISKS else ""}{f"| Naraj flow forecast, 1 / 2 / 3 days ahead (efficiency; 1 = perfect) | {fk['lead_1d']['boosted']['nse']:.2f} / {fk['lead_2d']['boosted']['nse']:.2f} / {fk['lead_3d']['boosted']['nse']:.2f} |" + chr(10) if FCS else ""}
## What it means for ClimMaTech

Their system lists a map-based River View layer, but no flood extent stands behind it; their alert areas are drawn
by hand; their solution pages show illustrative networks. This project adds the flood-extent and impact layers for one
delta, bulletins and CAP alerts for each forecast, a live feed that turns every day's forecast into an impact-based
warning, a network of sites placed with data, and a field method to turn their radar gauge readings into discharge.

## Limits and next steps

River flow is modelled (GloFAS), not gauged; pass-level skill is modest because floodplain flooding is patchy; depth
comes from a 30 m elevation model; soil, masonry and bed values are textbook values until the lab tests; the radar
misses shallow water among crops (it finds {s2f['recall']:.0%} of the flood water Sentinel-2 sees). Next: the field visit,
CWC/WRD flows, {"" if ADOPTED else f"the radar models retrained at 20 m (ready, {R20N.get('calibrated_test_iou', 0):.3f} against {R20.get('calibrated_test_iou', 0):.3f}), "}fine-tuning on the
delta with Sentinel-2 labels, and the Hirakud reservoir.

**Find everything:** report `report/FINAL_REPORT.docx`, slides `report/Prithvi_Flood-IN_Final.pptx`, website
`website/`, map viewer `dashboard/index.html`, live warnings `live/` (`python live/run_live.py`), field kit `fieldkit/`,
plan `PLAN.md`.
"""

# ---------------------------------------------------------------- speaker notes, keyed by the start of each slide title
NOTES = [
    ("Prithvi Flood-IN", "Nikhil", "Introduce the team and the question in one line: when the Mahanadi rises, what floods, and who must be warned? Say that the work joins satellite AI with the department's seven specializations, for a real delta and a real industry partner."),
    ("The problem in the field", "Nikhil", f"ClimMaTech measures and forecasts the river at points: radar gauges, sirens, an app. A forecast says how much water reaches Naraj, not which villages go under. Our refined problem statement fills that gap for the Mahanadi delta: {EXS['population_ghsl'] / 1e6:.1f} million people in 3,038 km²."),
    ("Who needs what", "Nikhil", "Walk through the stakeholders: ClimMaTech needs a flood-extent layer and site plans; the guide's Mahanadi project needs inundation and impacts; the Water Resources Department needs to know which banks and roads fail; the district needs villages to warn. Each row is something we deliver."),
    ("One flood, seven specializations", "Himanshu", "One set of flood scenarios, one question per specialization. Read the headline column quickly and say the details follow. Mention that water resources includes the watershed above the delta, where about half of a flood peak comes from below Hirakud dam."),
    ("Model selection: a fair benchmark", "Vineet", "Same data, same recipe, same decoder for every model: six foundation models, a CNN and a U-Net, three seeds each. The key design choice is leave-India-out: the test is an Indian flood no model saw. 92 runs in total."),
    ("Unseen Indian flood", "Vineet", f"The surprise: a small U-Net ({u('lio/unet/full/default'):.3f}) beat every foundation model (best {u('lio/prithvi/full/default'):.3f}). Water is spectrally simple and a few hundred labels are enough. Fusing the top three gives {u('lio/enstop3_mean/full/default'):.3f}."),
    ("Monsoon reality", "Vineet", "Optical is better under clear sky, but the delta was 61–95 % cloud on flood dates. So the production model is radar-only: TerraMind, which takes radar natively, fused with a U-Net."),
    ("Checked, not assumed", "Himanshu", f"We re-checked all 92 runs from their saved outputs: right settings, no chip in two splits, no Indian chip in training, selection on validation only, every score recomputed. Our TerraMind matches the published paper within a point ({om['official_terramind_full_s0']['miou'] * 100:.1f} against 90.62)."),
    ("Mahanadi flood atlas", "Prakhar", f"Every Sentinel-1 pass since 2017: {AS['n_passes']} passes, mapped in about 5 seconds each. Flood means water outside the river bed and the dry-season water. {n(AS['flooded_ever_km2'])} km² flooded at least once."),
    ("From a flow forecast to a flood map", "Nikhil", f"Each pixel gets the Naraj flow at which it starts to flood. Tested on years it never saw: IoU {HF['model']['iou']:.2f} against {HF['climatology']['iou']:.2f} for climatology at high flows. Past floods rank flood-prone land very well (AUC {auc['high_flow_passes']['clim']:.2f}); how much floods on a given day is harder."),
    ("A learned flow + rain + history model", "Nikhil", f"A gradient-boosted model adds 7-day rain and each pixel's flood history to the flow. It is better on heavy-rain passes ({LI['heavy_rain_passes']['learned']['iou']:.2f} against {LI['heavy_rain_passes']['trigger']['iou']:.2f}) and ranks land best (AUC {LA['all_passes']['learned']:.2f}), and it catches more flooded villages at the cost of more false alarms. Its maps never show less flooding for more rain."),
    ("Forecasting the Naraj flow", "Nikhil", f"How far ahead does the warning see? A model per lead time reads today's upstream flows and the catchment rain. On the test monsoons it beats assuming no change at every lead, most at three days ({fk.get('lead_3d', {}).get('boosted', {}).get('nse', 0):.2f} against {fk.get('lead_3d', {}).get('persistence', {}).get('nse', 0):.2f}). A rain forecast catches more 2-year floods two and three days ahead. With real gauges upstream the same code trains on measured flows."),
    ("What floods", "Himanshu", f"For each design flood: people on flooded land, farmland, roads, hospitals, bridges. In the 25-year flood {n(T25['people'])} people live on {T25['flood_km2']:.0f} km² of flooded land. Floods above the largest seen since 2017 are lower bounds."),
    ("Transportation", "Himanshu", f"The road network becomes a graph; flooded sections are cut. {n(TR25['hospital: people cut off'])} people lose road access to every hospital in the 25-year flood, and everyday trips grow by about {n(TR25['town: extra person-hours of travel per day'])} person-hours a day. The critical links are listed with coordinates."),
    ("Geotechnical: which river banks", "Prabhat", "Bank cells ranked by how often water stands behind the bank, the soil and the people nearby. Twenty inspection sites at least 3 km apart are listed, and the top twelve are on the field GPS file; the lab tests on their soil samples decide the ranking."),
    ("Structural", "Himanshu", f"Flood depth from the elevation at the flood edge. A 230 mm brick wall cracks at about a metre of water; {n(wT25['people_deeper_than_1m'])} people live where the 25-year flood is that deep. Pier scour reaches {scour25[0]:.0f}–{scour25[1]:.0f} m below the high flood level; the flow split between branches is the missing input."),
    ("Geotechnical: dry houses", "Prabhat", f"Floodwater raises the groundwater under houses that stay dry: {n(g25['people_in_zone'])} people within 200 m of the 25-year flood on silty clay. Soaking cuts bearing capacity by about {fe['silty_sand_drained']['reduction_pct']} %, and the clay settles when the water drops. Well levels at the site fix the 200 m."),
    ("Ocean", "Nikhil", f"{OCN['systems_within_300km_since_2016']} storms and depressions came within 300 km since 2016; {cnt['cyclone/depression']} of the 20 largest floods followed one. The two largest came at Naraj flows of "
     f"{' and '.join(n(t_.naraj_m3s) for t_ in TOPS)} m³/s, {'below' if all(t_.naraj_m3s < RP[2] for t_ in TOPS) else 'near'} the 2-year flood ({n(RP[2])}), with "
     f"{' and '.join(str(t_.rain_mm) for t_ in TOPS)} mm of rain that week: read the rain forecast with the river forecast."),
    ("Construction management", "Prabhat", f"Build in {win}, the months with no flood days and few storms. Raising {lr['links_by_embankment']} critical roads takes about {n(lr['earthwork_m3'])} m³ of earthwork{STRUCT}. The schedule fits approvals, procurement and installation inside the window."),
    ("Flood risk in rupees", "Himanshu", "Depth times a published damage curve gives the houses damaged; flooded paddy at the support price gives the crop lost. Running every day since 1997 through the impact table gives an average year: at least the expected damage on the slide, and the person-days without a hospital. Raising ten roads repays its cost many times over; the field interviews check the losses."),
    ("Deployment plan", "Prabhat", f"Sirens where exposed people live: {cov15['sirens']} sirens reach {cov15['share']:.0%} at 1.5 km, phones reach the rest. One radar gauge per distributary, on the bridge nearest Naraj whose approaches stay dry, measures how the flood splits."),
    ("Implementation", "Himanshu", f"A forecast flow goes in; a flood map, the settlements to warn and the roads under water come out in seconds, plus a CAP alert the Rakshak Control Centre can take ({B25.get('settlements_to_warn', 0)} alert areas for the 25-year flow). The map viewer shows every scenario."),
    ("Live warnings", "Himanshu", "Show today's run: 51 GloFAS runs become flows at Naraj, the flow–impact table turns them into people, villages and hospital access, and the level comes from the impact. Point at the villages with their first flood day, the official alerts the feed picked up, and the CAP alert and Atom feed the control centre can subscribe to. It runs in about 15 seconds."),
    ("Replay", "Himanshu", "Today is quiet, so here is the feed on the morning of the largest flood of the radar era. With the flows that came it says 'Act now' four days before the crest and names the villages with dates. The upstream model, which never saw 2022, under-forecasts the steep rise but still says 'Act now' three days ahead. The radar table shows the honest part: rain flooded more land before the crest than the crest itself."),
    ("The same pipeline on two more basins", "Prakhar", "A basin is a box, a gauge on the main river and the flood season. The same radar models mapped every pass since 2017 on the Brahmani-Baitarani delta and on the Brahmaputra at Majuli, the region of our unseen Indian test flood, and both run in the live feed."),
    ("What ClimMaTech's system lacks", "Himanshu", "We read every page of their site. Each row is one of their services, the gap and what we add, with the number behind it: the flood-extent layer for River View, bulletins and CAP alerts, a real network of sites, a field method for discharge, rain-aware maps."),
    ("Field visit", "Nikhil", "The visit measures what the models cannot know: accuracy at 40 random points, discharge and roughness at a bridge, bridge and bank inspections with soil samples, well levels, sensor sites and interviews. Every measurement has a script and changes a number in the report."),
    ("Checks of our own numbers", "Prakhar", f"Sentinel-2 within a day of 60 radar passes: IoU {s2a['iou']:.2f} on all water, lower on the flood part, so the radar misses shallow water among crops. The mentor's PINN for the Saint-Venant equations was reproduced exactly and retrained."),
    ("Team and contributions", "Nikhil", "Who built what: each member presents the part they own and can answer for the whole chain."),
    ("What we conclude", "Nikhil", "Seven points, one sentence each: the U-Net beats the foundation models on an unseen flood; the radar atlas turns a flow forecast into a map that beats climatology; rain matters, and the learned model uses it; the 25-year flood's people and roads, with an answer for every specialization; half of a flood peak comes from below Hirakud dam; every training run was re-checked; the tools are ready and the field visit closes the loop. Thank the guide and the mentors."),
]


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0]
    return ""


prs = Presentation(DECK)
rows = []
for i, sl in enumerate(prs.slides, 1):
    t = title_of(sl)
    match = next(((p, txt) for k, p, txt in NOTES if t.startswith(k)), ("—", ""))
    sl.notes_slide.notes_text_frame.text = (f"Presenter: {match[0]}. " + match[1]).strip()
    body, tables = [], []
    for sh in sl.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.text.split("\n"):
                para = para.strip().lstrip("•").strip()
                if para and para != t and not para.startswith("Prithvi Flood-IN ·"):  # skip the footer line
                    body.append(para)
        if sh.has_table:
            tb = [[c.text.replace("\n", " ").replace("|", "/") for c in r.cells] for r in sh.table.rows]
            tables.append("| " + " | ".join(tb[0]) + " |\n|" + "---|" * len(tb[0]) + "\n" +
                          "\n".join("| " + " | ".join(r) + " |" for r in tb[1:]))
    rows.append((i, t, match[0], body, tables, match[1]))
prs.save(DECK)
st = ["# Prithvi Flood-IN: slide text and speaker notes\n",
      "Every slide of `report/Prithvi_Flood-IN_Final.pptx`: what it shows, who presents it and what to say (about a minute per "
      "slide; the same notes are in the deck's speaker notes).\n"]
for i, t, who, body, tables, note in rows:
    st.append(f"## Slide {i}. {t}\n\n*Presenter: {who}*\n")
    if body:
        st.append("**On the slide**\n\n" + "\n".join(f"- {b}" for b in body) + "\n")
    for tb in tables:
        st.append(tb + "\n")
    if note:
        st.append(f"**What to say**\n\n{note}\n")
SLIDES = "\n".join(st)

# ---------------------------------------------------------------- the project from the basics
GUIDE = f"""# Prithvi Flood-IN: the whole project from the basics

For a team member who was not part of the work. Read it once from top to bottom (about 30 minutes), then read the
report sections of the part you will present. Every number here comes from the same result files as the report.

## 1. The project in one minute

When the Mahanadi rises, which fields, villages, roads, bridges and river banks in its delta go under water, and who
must be warned? Flood warning systems, including our industry partner ClimMaTech's, forecast the river level at a gauge
but cannot answer that. We (1) chose the best satellite model for mapping flood water with a fair test of eight
models, (2) used it to map every radar image of the delta since 2017, (3) linked those maps to the river flow at Naraj
so that a flow forecast becomes a flood map, (4) worked out what each flood means for each civil-engineering
specialization, and (5) built tools for the partner and a field visit that checks all of it on the ground.

## 2. Why the project looks like this

- **The IDP rules.** At the review on 14 October the coordinator asked for a real project with real stakeholders, a
  site visit whose measurements change the work, and input from several specializations. A team last year lost marks
  for a site visit that measured nothing.
- **The partner.** ClimMaTech (an IIT Mandi start-up) makes radar river gauges (Pravaah, Tarang), a camera unit
  (Doordrishti), solar sirens (Rakshak), rain gauges (Varshaq), a control centre and the Aagah app. Every one of these
  works at a point. Their own website lists a map layer for flood risk (River View) but nothing that shows flood extent.
- **The guide.** Prof. Basudev Biswal works on the Mahanadi basin. The Mahanadi delta below Naraj (around Cuttack and
  Kendrapara) floods every few years: a 3,038 km² study box holds about {EXS['population_ghsl'] / 1e6:.1f} million people.

## 3. Background you need

### 3.1 Two kinds of satellite image

**Sentinel-2** is an optical satellite, a camera with extra colours (13 bands). Water is easy to see in it, but clouds
block it, and the Mahanadi floods in the monsoon: on the big flood dates the delta was 61–95 % cloud. **Sentinel-1**
is a radar satellite. It sends microwaves and records the echo, so it sees through cloud, day and night. Calm water
reflects the radar away like a mirror and looks dark; land scatters it back and looks brighter. It passes over the
delta every 2–6 days. That is why the flood atlas uses radar.

### 3.2 What "mapping water" means for a model

A model takes an image tile (512 × 512 pixels, about 5 km × 5 km at 10 m) and gives every pixel a probability of being
water. Above a threshold (usually 0.5) the pixel is called water. This is called segmentation. The model learns from
images where people have marked the water by hand (the labels).

### 3.3 Foundation models and small networks

A **foundation model** is a large network (35 to 324 million parameters) pre-trained on millions of unlabelled
satellite images; it is then fine-tuned for a task. We tested six: Prithvi-EO-2.0 (IBM/NASA), TerraMind (IBM/ESA),
DOFA, Clay, Satlas and DeCUR. A **U-Net** (7.8 million parameters) and a three-layer **CNN** are small networks
trained from scratch, the baselines. **Fine-tuned** means all weights keep learning; **frozen** means only a small
decoder learns on top of the pre-trained features. **Fusion** averages or combines the predictions of several models.

### 3.4 How we score

- **IoU** (intersection over union) = overlap of predicted and true water ÷ their union; 1 is perfect. A water IoU of
  0.77 is good for this dataset.
- **Precision**: of the pixels we called water, the share that is water. **Recall**: of the true water, the share we found.
- **POD / FAR**: for villages, the share of flooded villages we flagged, and the share of our flags that were wrong.
- **ROC AUC**: how well a map ranks flooded land above dry land (0.5 = chance, 1 = perfect).

## 4. Part 1: choosing the model

**Data.** Sen1Floods11: 446 hand-labelled tiles from 11 floods worldwide, each with Sentinel-1 and Sentinel-2.
**Two tests.** The *official* split, comparable with published papers, and *leave-India-out*: train without any Indian
tile and test on all 68 tiles of the 2016 Assam flood, a flood no model saw. The second is the honest test for India.

**What we trained (92 runs, three seeds for most):**

| Group | Runs | Question |
|---|---|---|
| Six foundation models, fully fine-tuned | 36 | Which one maps flood water best? |
| CNN and U-Net from scratch | 12 | Do foundation models beat small networks? |
| Six foundation models, frozen | 12 | How much do the pre-trained features know? |
| Radar only or optical only | 14 | Which satellite matters? |
| 10, 25 and 50 % of the labels | 18 | Do foundation models need fewer labels? |

Every model used the same recipe (same loss, optimiser, augmentation, 40 epochs) and the epoch with the best
validation score; the test tiles were used once, at the end.

**Results.** On the unseen Assam flood the U-Net scored {u('lio/unet/full/default'):.3f}, the best foundation model
(Prithvi) {u('lio/prithvi/full/default'):.3f}, and a fusion of the U-Net, the CNN and the best foundation model
{u('lio/enstop3_mean/full/default'):.3f}. The small network won because water is spectrally simple and a few hundred labelled
tiles are enough; the U-Net stayed ahead even with 10 % of the labels. Radar alone scores lower (about 0.65 on the
official test, against 0.82–0.84 with radar and optical together), but it is the only thing that sees a monsoon
flood, and TerraMind takes radar natively. **The model we use** is therefore the radar-only TerraMind + U-Net fusion.
{f"Retrained directly on 20 m input it scores {R20N.get('calibrated_test_iou', 0):.3f} IoU on the 20 m radar the atlas reads (the pair trained at 10 m: {R20['calibrated_test_iou']:.3f}), so the atlas was rebuilt with it." if ADOPTED else f"It scores {R20['calibrated_test_iou']:.3f} IoU on the 20 m radar the atlas reads; retrained directly at 20 m the pair reaches {R20N.get('calibrated_test_iou', 0):.3f}, ready for the next atlas run."}

**Was the training right?** A script re-checked all 92 runs from their saved outputs: correct settings, no tile in
two splits, no Indian tile in leave-India-out training, selection on validation only, every score recomputed: no
problems. Our TerraMind on the official test scores {om['official_terramind_full_s0']['miou'] * 100:.1f} mean IoU (radar +
optical) and {om['official_terramind_s1_full_s0']['miou'] * 100:.1f} (radar only); the published paper reports 90.62 and 80.39.

## 5. Part 2: the flood atlas

The model mapped every Sentinel-1 pass over the delta from 2017 to October 2026: {AS['n_passes']} passes, {AS['n_monsoon']} in the
monsoon (June–November) and {AS['n_dry']} in the dry season (February–March), at 20 m, about 5 seconds per pass on one GPU.
Water in the dry season and the river bed ({AS['river_bed_km2']:.0f} km²) are "normal water"; water outside them on a monsoon
pass is **flood**. {n(AS['flooded_ever_km2'])} km² flooded at least once. **Independent check:** for 60 passes a Sentinel-2
image exists within a day. On clear pixels the two agree with IoU {s2a['iou']:.2f} on all water, but the radar finds only
{s2f['recall']:.0%} of the flood water Sentinel-2 sees: it misses shallow water among crops. Field points will check it again.

## 6. Part 3: from a river flow to a flood map

**River flow** at Naraj, the head of the delta, comes from GloFAS, a global model (daily, 1997–2026); the measured
CWC/WRD flows are requested. The flooded share of a pass follows {PRED_WORDS} (rank correlation
{AS['predictor_spearman'][AS['predictor']]:.2f}).

**Trigger discharge.** For each 20 m pixel we find the Naraj flow above which it is usually flooded. A forecast flow
then gives a map: every pixel whose trigger is below the forecast. **Tested on years it never saw** (each year
predicted from the other nine): IoU {ALLP['model']['iou']:.2f} over all monsoon passes against {ALLP['climatology']['iou']:.2f} for
"flood where it usually floods", and {HF['model']['iou']:.2f} against {HF['climatology']['iou']:.2f} at high flows. Per village it flags
{sett_hf['POD']:.0%} of flooded villages. These scores are modest because floodplain flooding is patchy, and the honest
summary is: *where* the delta floods is very predictable from its past (ranking AUC {auc['high_flow_passes']['clim']:.2f}),
*how much* floods on a given day only partly.

**Rain matters too.** The largest floods ({" and ".join(f"{t_['when']:%B %Y}" for t_ in TOPS)}) came with heavy rain{" from Bay of Bengal depressions" if all(isinstance(t_.system, str) and t_.system for t_ in TOPS) else ""},
not with the highest river flows. At the 5-year flow, a dry week floods {RB[20]['flooded_land_km2']:.0f} km² and a
heavy-rain week {RB[RH]['flooded_land_km2']:.0f} km² (past passes with similar flow and rain). A **learned model**
(gradient boosting) that adds rain and each pixel's flood history to the flow does best on heavy-rain passes
({LI['heavy_rain_passes']['learned']['iou']:.2f} against {LI['heavy_rain_passes']['trigger']['iou']:.2f}) and ranks land best (AUC {LA['all_passes']['learned']:.2f}).

{f"**Forecasting the flow.** A model reads today's flows at Naraj, Tikarpara and below Hirakud and the rain on the catchment between them, and forecasts the Naraj flow one to three days ahead. On monsoons it never saw it scores an efficiency of {fk['lead_1d']['boosted']['nse']:.2f}, {fk['lead_2d']['boosted']['nse']:.2f} and {fk['lead_3d']['boosted']['nse']:.2f} (1 = perfect) against {fk['lead_1d']['persistence']['nse']:.2f}, {fk['lead_2d']['persistence']['nse']:.2f} and {fk['lead_3d']['persistence']['nse']:.2f} for assuming the flow stays as it is." + chr(10) + chr(10) if FCS else ""}**Design floods.** A Gumbel fit to {AS['gumbel']['n_years']} yearly maxima (1997–2025) gives {n(RP[2])} m³/s for the 2-year flood, {n(RP[5])} for the 5-year,
{n(RP[25])} for the 25-year and {n(RP[100])} for the 100-year. The largest flow in the radar record is
{n(AS['largest_observed_predictor_m3s'])} m³/s, so maps for the 10- to 100-year floods are the same, a lower bound.

## 7. Part 4: the specializations, briefly

- **Transportation:** roads as a network; in the 25-year flood {n(TR25['hospital: people cut off'])} people lose road access to
  every hospital and everyday trips grow by {n(TR25['town: extra person-hours of travel per day'])} person-hours a day; ten critical
  roads carry most of the cut routes.
- **Structural:** a 230 mm brick wall cracks at about a metre of water, and {n(wT25['people_deeper_than_1m'])} people live where
  the 25-year flood is that deep; bridge piers may scour {scour25[0]:.0f}–{scour25[1]:.0f} m below the flood level.
- **Geotechnical:** twenty candidate river-bank sites, the top twelve on the field GPS file; {n(g25['people_in_zone'])} people in dry houses within 200 m of the
  flood on soft soil, where foundations weaken and later settle.
- **Ocean:** {cnt['cyclone/depression']} of the 20 largest floods followed a Bay of Bengal storm.
- **Construction management:** build in {win}; raising {lr['links_by_embankment']} critical roads takes about
  {n(lr['earthwork_m3'])} m³ of earthwork; a schedule for sirens and gauges.
- **Water resources and the watershed:** Hirakud dam controls about 59 % of the basin, but only about half of a
  flood peak at Naraj (median {WS['median_share_hirakud_of_naraj_peak']:.0%}) comes from it; the rest comes from the tributaries
  below the dam within a day.
- **Remote sensing:** Parts 1 and 2.
{f"- **Risk in rupees:** flood depth times a published damage curve for houses, and flooded paddy at the support price: about ₹{RISKS['expected_annual_damage_rs_crore']['total_rs_crore']:.0f} crore of damage a year (a lower bound). Raising the ten critical roads repays its cost about {RISKS['road_raising']['benefit_cost_ratio']:.0f} times in hospital access alone." + chr(10) if RISKS else ""}
## 8. Part 5: the tools

- **Bulletin tool** (`mahanadi/forecast_to_impact.py`): type a forecast flow (m³/s or lakh cusec, optionally rain);
  get a flood map, the settlements to warn, the roads under water and a CAP alert, the format ClimMaTech's control
  centre accepts. For the 25-year flow it names {B25.get('settlements_to_warn', 0)} settlements.
- **Map viewer** (`dashboard/index.html`): every scenario with roads, sirens, gauges and banks; works offline.
- **Siting:** {cov15['sirens']} sirens reach {cov15['share']:.0%} of exposed people at 1.5 km; one radar gauge on each distributary.
- **Website** (`website/`): the whole project in the partner's visual style.
{f"- **Live warnings** (`live/run_live.py`, about 15 seconds): reads today's GloFAS ensemble, ECMWF rain and the official alerts and writes the warning level, the villages to warn with dates, a CAP alert and an Atom feed. Levels are set by impact: 'Be aware' from {n(LCFG['level_flows_m3s'].get('1'))} m³/s at Naraj, 'Be prepared' from {n(LCFG['level_flows_m3s'].get('2'))}, 'Act now' from {n(LCFG['level_flows_m3s'].get('3'))}. Each data source is a plugin (`live/sources.py`)." + REPLAY_TXT + chr(10) if LIVES and LCFG else ""}{"- **Other basins:** the same pipeline mapped the " + " and the ".join(BNAMES) + " from every radar pass since 2017" + LEVELS_TXT + "." + chr(10) if BASK else ""}
## 9. Part 6: the field visit

Three days in the delta, with protocols A–H in `fieldkit/FIELD_VISIT_KIT.md`: 40 random ground-truth points (did this
spot flood, how deep), discharge and roughness at a road bridge, bridge and river-bank inspections with soil samples
for the IIT Bombay lab, water levels in village wells, siren and gauge sites, and interviews with WRD, CWC, ClimMaTech
and villagers. Each measurement has a script that turns it into a number in report chapter 8. Load
`fieldkit/checkpoints/checkpoints.gpx` on a phone before going.

## 10. Honest limits

River flow is modelled, not gauged. Depth comes from a 30 m elevation model (about 2 m error). Soil, masonry and
riverbed values are textbook values until the lab tests. The flow split between the delta's branches is unknown.
The radar misses shallow water among crops. Floods larger than any since 2017 are lower bounds. Say these plainly if
asked: each one has a next step.

## 11. Questions you may be asked

- **Why not just use the published Prithvi flood model?** It was trained on the official split, which includes Indian
  tiles, so its Indian score is not an unseen test; on our unseen test it scored {u('lio/prithvi_published/none/default'):.3f}, below the U-Net.
- **If the U-Net won, why use a foundation model at all?** For the monsoon only radar works, and TerraMind takes radar
  natively; the radar fusion is the model we deploy.
- **How do you know the atlas is right?** Sentinel-2 on the same days, years held out from fitting, and the field points.
- **Why is the pass-level score only about 0.2?** Floodplain flooding is patchy, the observed maps are themselves
  radar maps, and rain drives the largest floods; the frequency map ranks land with AUC above 0.9.
- **Why are the 10- to 100-year maps identical?** The radar has not yet seen a flood larger than about a 10-year one.
- **Did you check the training?** Yes: 92 of 92 runs re-checked, no problems, published TerraMind scores reproduced.
- **What does ClimMaTech get?** A flood-extent layer, bulletins with CAP alerts, a network of sites placed with data,
  and a field method to turn gauge readings into discharge.
- **Is the live feed an official warning?** No. Its CAP alerts carry the status "Exercise"; it shows what an
  impact-based service would say, for ClimMaTech and the district to adopt.
- **Why is a village that floods every monsoon not in the warning?** Levels count what a flood adds to the ordinary
  monsoon (the flow exceeded on a quarter of monsoon days{f", {n(LCFG['baseline_m3s'])} m³/s at Naraj" if LCFG.get('baseline_m3s') else ""}). A few
  river-bank hamlets are wet at that flow every year; on the Brahmani a road with no mapped bridge cuts six villages off
  at almost any monsoon flow, which is a road to raise, not a warning to issue.
- **How far ahead can you warn?** GloFAS gives 30 days; our upstream-gauge model is good to about three days
  (efficiency {fk.get('lead_3d', {}).get('boosted', {}).get('nse', 0):.2f}); the villages list says the first day each is expected to flood.
- **What is next?** The field visit, observed flows, {'' if ADOPTED else 'the 20 m models in the atlas, '}Sentinel-2 labels for the delta.

## 12. Where everything is

| What | Where |
|---|---|
| The plan, roles and timeline | `PLAN.md` |
| Final report | `report/FINAL_REPORT.docx` (Appendix E is a glossary) |
| Slides, with speaker notes | `report/Prithvi_Flood-IN_Final.pptx`; text in `docs/SLIDES_TEXT.docx` |
| Executive summary | `docs/EXECUTIVE_SUMMARY.docx` |
| Field visit kit and GPS points | `fieldkit/` |
| Map viewer and website | `dashboard/index.html`, `website/` |
| Live warnings (run it before the presentation) | `live/` — `python live/run_live.py`, then open `live/out/index.html`; `live/README.md` |
| How to re-run anything | `README.md` |
"""

for name, text in (("EXECUTIVE_SUMMARY", EXEC), ("SLIDES_TEXT", SLIDES), ("TEAM_GUIDE", GUIDE)):
    md = os.path.join(DOCS, f"{name}.md")
    open(md, "w", encoding="utf-8").write(text)
    pypandoc.convert_file(md, "docx", outputfile=os.path.join(DOCS, f"{name}.docx"))
    print(f"written docs/{name}.md and .docx", len(text.split()), "words")
print("speaker notes written into", os.path.basename(DECK))
