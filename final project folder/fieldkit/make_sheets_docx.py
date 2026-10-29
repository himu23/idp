"""Printable field data sheets (Field_data_sheets.docx) from the CSV templates, and Word copies of the kit and plan.

    python fieldkit/make_sheets_docx.py
"""
import csv
import os

import docx
import pypandoc
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TITLES = {
    "ground_truth_points": ("A. Ground-truth points and flood marks", 14),
    "cross_section": ("B. River cross-section (from the bridge)", 16),
    "float_velocity": ("C. Float timings", 16),
    "water_surface_slope": ("C. Water-surface slope", 4),
    "manning_cowan": ("C. Manning's n by Cowan's method (words from the table in analysis/manning_cowan.py)", 4),
    "bridge_inspection": ("D. Bridge inspection (ask the foundation level)", 6),
    "embankment_inspection": ("E. Embankment / river-bank inspection", 8),
    "soil_samples": ("E. Soil samples", 10),
    "sensor_site_assessment": ("F. Sensor / siren site assessment", 10),
    "interviews": ("G. Interviews", 8),
    "well_levels": ("H. Water level in village wells (distance from the flood edge from the dossier map)", 12),
    "photo_log": ("Photo log", 18),
}
d = docx.Document()
sec = d.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(1.2))
d.add_heading("Mahanadi delta field visit — data sheets", 0)
d.add_paragraph("Fill on paper, type into the CSV of the same name in fieldkit/templates the same evening. "
                "Every photo gets a row in the photo log. Protocols: fieldkit/FIELD_VISIT_KIT.md.")
for name, (title, rows) in TITLES.items():
    cols = next(csv.reader(open(os.path.join(HERE, "templates", f"{name}.csv"), encoding="utf-8")))
    d.add_page_break()
    d.add_heading(title, 1)
    d.add_paragraph(f"File: templates/{name}.csv   Date: ________   Recorder: ____________   Site: ____________")
    t = d.add_table(rows=rows + 1, cols=len(cols))
    t.style = "Table Grid"
    for j, c in enumerate(cols):
        cell = t.cell(0, j)
        cell.text = c.replace("_", " ")
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(7 if len(cols) > 10 else 8)
                r.font.bold = True
    for i in range(1, rows + 1):
        t.rows[i].height = Cm(0.8)
d.save(os.path.join(HERE, "Field_data_sheets.docx"))
root = os.path.dirname(HERE)
pypandoc.convert_file(os.path.join(HERE, "FIELD_VISIT_KIT.md"), "docx", outputfile=os.path.join(HERE, "FIELD_VISIT_KIT.docx"))
pypandoc.convert_file(os.path.join(root, "PLAN.md"), "docx", outputfile=os.path.join(root, "PLAN.docx"))
print("written Field_data_sheets.docx,", len(TITLES), "sheets; FIELD_VISIT_KIT.docx; PLAN.docx")
